#!/usr/bin/env python3
"""Read-only PPTX inventory. Explicit XML formatting is not the resolved rendered style."""
import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path
from zipfile import BadZipFile

try:
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError
    from pptx.oxml.ns import qn
    from lxml.etree import XMLSyntaxError
except ImportError as error:
    raise SystemExit("Install dependency: python -m pip install 'python-pptx>=1.0.2'") from error

INHERITED = "<inherited>"
STYLE_FIELDS = ("font", "size", "weight", "color", "font_east_asian", "font_complex_script", "italic")
REGIONS = ("title_header", "upper_explanation", "middle_visual", "lower_conclusion", "unresolved")


def child(element, tag):
    return element.find(qn(tag)) if element is not None else None


def fill_value(properties):
    if properties is None:
        return None
    solid = child(properties, "a:solidFill")
    if solid is not None and len(solid):
        color = solid[0]
        kind = color.tag.rsplit("}", 1)[-1]
        val = color.get("val", "")
        prefixes = {"srgbClr": "#", "schemeClr": "theme:", "sysClr": "system:"}
        result = prefixes.get(kind, kind + ":") + val
        if kind in ("hslClr", "scrgbClr"):
            result = kind + ":" + ",".join(f"{key}={value}" for key, value in sorted(color.attrib.items()))
        if kind == "sysClr":
            result += ";last=" + color.get("lastClr", "")
        for modifier in color:
            result += ";" + modifier.tag.rsplit("}", 1)[-1] + "=" + modifier.get("val", "")
        return result
    for kind in ("noFill", "gradFill", "pattFill", "blipFill", "grpFill"):
        if child(properties, "a:" + kind) is not None:
            return kind
    return None


def font_style(properties):
    def attr(name):
        return properties.get(name) if properties is not None else None

    def face(tag):
        node = child(properties, tag)
        return (node.get("typeface") or INHERITED) if node is not None else INHERITED

    def boolean(name):
        value = attr(name)
        return INHERITED if value is None else value in ("1", "true", "on")

    bold, size = boolean("b"), attr("sz")
    return {
        "font": face("a:latin"),
        "size": f"{int(size) / 100:g}pt" if size is not None else INHERITED,
        "weight": INHERITED if bold == INHERITED else "bold" if bold else "regular",
        "color": fill_value(properties) or INHERITED,
        "font_east_asian": face("a:ea"),
        "font_complex_script": face("a:cs"),
        "italic": boolean("i"),
    }


def frame_styles(text_frame):
    counts = Counter()
    for paragraph in text_frame.paragraphs:
        # Read XML directly: font.color and get_or_add APIs can mutate the in-memory XML.
        for run in paragraph._p:
            if run.tag not in (qn("a:r"), qn("a:fld")):
                continue
            node = child(run, "a:t")
            text = (node.text or "").strip() if node is not None else ""
            if text:
                style = font_style(child(run, "a:rPr"))
                counts[tuple(style[field] for field in STYLE_FIELDS)] += len(text)
    return counts


def text_style(shape):
    return frame_styles(shape.text_frame) if shape.has_text_frame else Counter()


def style_records(counts, limit=None):
    return [dict(zip(STYLE_FIELDS, style), characters=count)
            for style, count in counts.most_common(limit)]


def paragraph_defaults(text_frame):
    result = []
    for number, paragraph in enumerate(text_frame.paragraphs, 1):
        props = child(paragraph._p, "a:pPr")
        result.append({
            "paragraph": number,
            "level": int(props.get("lvl", "0")) if props is not None else 0,
            "alignment": props.get("algn", INHERITED) if props is not None else INHERITED,
            "explicit_text_style": font_style(child(props, "a:defRPr")),
        })
    return result


def parse_slide_numbers(raw, count):
    """Parse 1-based numbers/ranges, preserving first occurrence order."""
    if raw is None or raw == "":
        return []
    numbers = []
    for token in raw.split(","):
        match = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d+)\s*)?", token)
        if not match:
            raise ValueError(f"Invalid slide selection: {token!r}; use 1,3-5")
        start = int(match[1])
        end = int(match[2]) if match[2] else start
        if start < 1 or end > count or start > end:
            raise ValueError(f"Slide range {token.strip()!r} is outside 1-{count} or reversed")
        for number in range(start, end + 1):
            if number not in numbers:
                numbers.append(number)
    return numbers


def shape_box(shape, transform):
    if transform is None or any(v is None for v in (shape.left, shape.top, shape.width, shape.height)):
        return None
    sx, sy, tx, ty = transform
    xs = (shape.left * sx + tx, (shape.left + shape.width) * sx + tx)
    ys = (shape.top * sy + ty, (shape.top + shape.height) * sy + ty)
    return min(xs), min(ys), abs(xs[1] - xs[0]), abs(ys[1] - ys[0])


def norm_box(box, width, height):
    return {key: round(value / axis * 100, 2) for key, value, axis in zip(
        ("left_pct", "top_pct", "width_pct", "height_pct"), box,
        (width, height, width, height))} if box is not None else None


def zone(box, height):
    if box is None:
        return "unresolved"
    top = box[1] / height
    if top < 0.12:
        return "title_header"
    if top < 0.30:
        return "upper_explanation"
    return "lower_conclusion" if top >= 0.82 else "middle_visual"


def group_transform(shape, parent):
    """Map unrotated child coordinates through group scale/translation/flips."""
    xfrm = shape._element.xfrm
    if xfrm is None or xfrm.chOff is None or xfrm.chExt is None:
        return None
    if not xfrm.chExt.cx or not xfrm.chExt.cy:
        return None
    sx, sy = shape.width / xfrm.chExt.cx, shape.height / xfrm.chExt.cy
    if xfrm.flipH:
        sx = -sx
    if xfrm.flipV:
        sy = -sy
    tx = shape.left + (shape.width if xfrm.flipH else 0) - xfrm.chOff.x * sx
    ty = shape.top + (shape.height if xfrm.flipV else 0) - xfrm.chOff.y * sy
    psx, psy, ptx, pty = parent
    return sx * psx, sy * psy, tx * psx + ptx, ty * psy + pty


def walk_shapes(shapes, transform=(1, 1, 0, 0), parent_path=(), status="slide"):
    for shape in shapes:
        path = parent_path + (shape.shape_id,)
        is_group = shape._element.tag == qn("p:grpSp")
        current_status = status
        if status == "slide" and is_group and shape.rotation % 360:
            current_status = "unsupported_group_rotation"
        box = shape_box(shape, transform) if current_status == "slide" else None
        if box is None and current_status == "slide":
            current_status = "missing_geometry"
        yield shape, path, parent_path, box, current_status
        if is_group:
            next_transform = group_transform(shape, transform) if current_status == "slide" else None
            child_status = current_status
            if current_status == "slide" and next_transform is None:
                child_status = "unsupported_group_extent"
            yield from walk_shapes(shape.shapes, next_transform, path, child_status)


def element_details(shape, path, parent_path, box, status, width, height):
    result = {
        "name": shape.name, "type": str(shape.shape_type),
        "path": list(path), "parent_path": list(parent_path),
        "region": zone(box, height), "box": norm_box(box, width, height),
        "coordinate_status": status, "rotation_degrees": shape.rotation,
    }
    if shape.has_text_frame and shape.text.strip():
        result.update(text=shape.text.strip()[:180],
                      explicit_text_styles=style_records(text_style(shape)),
                      paragraph_defaults=paragraph_defaults(shape.text_frame))
    if shape.is_placeholder:
        result["placeholder"] = {"idx": shape.placeholder_format.idx,
                                 "type": str(shape.placeholder_format.type)}
    props = child(shape._element, "p:spPr")
    fill = fill_value(props)
    if fill is not None:
        result["fill"] = fill
    line = child(props, "a:ln")
    if line is not None:
        result["line"] = {"color": fill_value(line) or INHERITED,
                          "width_pt": int(line.get("w")) / 12700 if line.get("w") else INHERITED}
    if shape.has_table:
        result["table_cells"] = [
            {"row": row_no, "column": col_no, "text": cell.text[:180],
             "is_merge_origin": cell.is_merge_origin, "is_spanned": cell.is_spanned,
             "explicit_text_styles": style_records(frame_styles(cell.text_frame)),
             "paragraph_defaults": paragraph_defaults(cell.text_frame)}
            for row_no, row in enumerate(shape.table.rows, 1)
            for col_no, cell in enumerate(row.cells, 1)]
    return result


def inspect_presentation(path, slides=None):
    """Summarize all slides; aggregate selected slides and detail them in order.

    slides=None aggregates all slides without detailed elements. No source is saved.
    """
    path = Path(path)
    deck = Presentation(str(path))
    width, height = deck.slide_width, deck.slide_height
    if not width or not height:
        raise ValueError("Presentation must have a nonzero slide width and height")
    selected = list(dict.fromkeys(slides)) if slides is not None else []
    if any(not isinstance(n, int) or isinstance(n, bool) or n < 1 or n > len(deck.slides) for n in selected):
        raise ValueError(f"Slide numbers must be integers in 1-{len(deck.slides)}")
    scope = selected if slides is not None else list(range(1, len(deck.slides) + 1))
    region_styles = {region: Counter() for region in REGIONS}
    summaries, details = [], {}
    for index, slide in enumerate(deck.slides, 1):
        entries = list(walk_shapes(slide.shapes))
        summaries.append({
            "slide": index, "layout": slide.slide_layout.name,
            "shape_count": len(slide.shapes), "total_shape_count": len(entries),
            "text_boxes": sum(bool(s.has_text_frame and s.text.strip()) for s, *_ in entries),
            "images": sum(s._element.tag == qn("p:pic") for s, *_ in entries),
            "tables": sum(s.has_table for s, *_ in entries),
            "charts": sum(s.has_chart for s, *_ in entries),
        })
        if index in scope:
            for shape, _, _, box, _ in entries:
                counts = text_style(shape)
                if shape.has_table:
                    for cell in shape.table.iter_cells():
                        if not cell.is_spanned:
                            counts.update(frame_styles(cell.text_frame))
                region_styles[zone(box, height)].update(counts)
        if index in selected:
            details[index] = {"slide": index, "elements": [
                element_details(*entry, width, height) for entry in entries]}
    return {
        "format_version": 2, "source": str(path.resolve()),
        "slide_count": len(deck.slides),
        "slide_size_inches": [round(width / 914400, 4), round(height / 914400, 4)],
        "style_scope_slides": scope,
        "explicit_text_styles_by_region": {
            region: style_records(counts, 12) for region, counts in region_styles.items()},
        "slides": summaries,
        "selected_slide_details": [details[n] for n in selected],
        "caution": (
            "Explicit XML formatting only; <inherited> is unresolved, not a default font or size. "
            "Theme colors and font references are not resolved. Check paragraphs, placeholders, "
            "layouts, masters, theme and rendered slides. Regions are top-coordinate heuristics; "
            "they do not determine slide roles. Boxes are nominal unrotated extents. Rotated groups "
            "and missing group transforms have null boxes and require visual inspection. "
            "Table styles, chart text, SmartArt, master-only artwork and image text are not inventoried."
        ),
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path, help="Source PPTX; never modified")
    parser.add_argument("--slides", help="1-based numbers/ranges, e.g. 2,4-6; scope for styles and details")
    parser.add_argument("--output", type=Path, help="Write UTF-8 JSON instead of stdout; cannot be the source")
    args = parser.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if not args.pptx.is_file():
        parser.error(f"PPTX does not exist: {args.pptx}")
    try:
        if args.output is not None and (args.output.resolve() == args.pptx.resolve() or (
                args.output.exists() and args.output.samefile(args.pptx))):
            parser.error("Output must not overwrite the source PPTX")
        selected = None
        if args.slides is not None:
            count = len(Presentation(str(args.pptx)).slides)
            selected = parse_slide_numbers(args.slides, count)
            if not selected:
                raise ValueError("--slides must contain at least one slide number")
        report = inspect_presentation(args.pptx, selected)
        encoded = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
        if args.output is not None:
            args.output.write_text(encoded, encoding="utf-8")
        else:
            sys.stdout.write(encoded)
    except (OSError, ValueError, KeyError, BadZipFile, PackageNotFoundError, XMLSyntaxError) as error:
        parser.error(f"Cannot inspect PPTX: {error}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

