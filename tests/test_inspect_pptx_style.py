"""Regression tests using generated OOXML fixtures, without private slide decks.

Run with: python -m unittest discover -s tests -v
If the system temp folder is restricted, set SANDWICH_TEST_TMPDIR to a writable
folder such as _local/tmp. All generated decks are removed after each test.
"""

import base64
import hashlib
import importlib.util
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "inspect_pptx_style.py"
SPEC = importlib.util.spec_from_file_location("sandwich_style_inspector", SCRIPT)
INSPECTOR = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = INSPECTOR
SPEC.loader.exec_module(INSPECTOR)


def add_text(shapes, name, text, left=1, top=2, width=3, height=1):
    shape = shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.name = name
    shape.text_frame.paragraphs[0].add_run().text = text
    return shape


def styles_for(report, name, slide=0):
    return next(
        element
        for element in report["selected_slide_details"][slide]["elements"]
        if element["name"] == name
    )


class SlideSelectionTests(unittest.TestCase):
    def test_empty_and_ordered_duplicate_selection(self):
        self.assertEqual(INSPECTOR.parse_slide_numbers(None, 5), [])
        self.assertEqual(INSPECTOR.parse_slide_numbers("", 5), [])
        self.assertEqual(INSPECTOR.parse_slide_numbers(" 3,1,3,2 ", 5), [3, 1, 2])

    def test_inclusive_ranges(self):
        self.assertEqual(INSPECTOR.parse_slide_numbers("3-5,1,4", 5), [3, 4, 5, 1])

    def test_invalid_selections_are_rejected(self):
        for raw in ("0", "-1", "6", "x", "1,,2", "3-1", "1-6", "1-2-3"):
            with self.subTest(raw=raw), self.assertRaises(ValueError):
                INSPECTOR.parse_slide_numbers(raw, 5)


class InspectorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(
            prefix="sandwich-ppt-test-", dir=os.environ.get("SANDWICH_TEST_TMPDIR")
        )
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "fixture.pptx"
        self.deck = Presentation()
        self.deck.slide_width = Inches(10)
        self.deck.slide_height = Inches(7.5)
        self.slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])

    def inspect(self, slides=None):
        self.deck.save(self.path)
        return INSPECTOR.inspect_presentation(self.path, slides=slides)

    def cli(self, *arguments, encoding=None):
        env = os.environ.copy()
        if encoding is not None:
            env["PYTHONIOENCODING"] = encoding
        return subprocess.run(
            [sys.executable, str(SCRIPT), *map(str, arguments)],
            capture_output=True,
            env=env,
            check=False,
        )

    def test_mixed_explicit_run_styles_remain_distinct(self):
        shape = add_text(self.slide.shapes, "Mixed", "Alpha")
        first = shape.text_frame.paragraphs[0].runs[0]
        first.font.name = "Arial"
        first.font.size = Pt(20)
        first.font.bold = True
        first.font.italic = False
        first.font.color.rgb = RGBColor.from_string("003399")
        second = shape.text_frame.paragraphs[0].add_run()
        second.text = "Beta"
        second.font.name = "Calibri"
        second.font.size = Pt(18)
        second.font.bold = False
        second.font.italic = True
        report = self.inspect([1])
        styles = styles_for(report, "Mixed")["explicit_text_styles"]
        self.assertEqual(len(styles), 2)
        alpha = next(style for style in styles if style["font"] == "Arial")
        beta = next(style for style in styles if style["font"] == "Calibri")
        self.assertEqual((alpha["size"], alpha["weight"], alpha["color"]),
                         ("20pt", "bold", "#003399"))
        self.assertEqual((beta["size"], beta["weight"]), ("18pt", "regular"))
        self.assertEqual((alpha["characters"], beta["characters"]), (5, 4))
        self.assertNotEqual(alpha["italic"], beta["italic"])

    def test_cjk_and_complex_script_fonts_are_not_lost(self):
        shape = add_text(self.slide.shapes, "Multilingual", "中文日本語한글")
        run = shape.text_frame.paragraphs[0].runs[0]
        run.font.name = "Arial"
        for tag, face in (("a:ea", "Microsoft YaHei"), ("a:cs", "Arial Unicode MS")):
            node = OxmlElement(tag)
            node.set("typeface", face)
            run._r.get_or_add_rPr().append(node)
        style = styles_for(self.inspect([1]), "Multilingual")["explicit_text_styles"][0]
        self.assertEqual(style["font"], "Arial")
        self.assertEqual(style["font_east_asian"], "Microsoft YaHei")
        self.assertEqual(style["font_complex_script"], "Arial Unicode MS")
        self.assertEqual(style["characters"], 7)

    def test_theme_color_modifiers_are_preserved(self):
        shape = add_text(self.slide.shapes, "Theme color", "Tinted text")
        properties = shape.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
        fill = OxmlElement("a:solidFill")
        color = OxmlElement("a:schemeClr")
        color.set("val", "accent1")
        for name, value in (("lumMod", "75000"), ("alpha", "50000")):
            modifier = OxmlElement("a:" + name)
            modifier.set("val", value)
            color.append(modifier)
        fill.append(color)
        properties.append(fill)
        style = styles_for(self.inspect([1]), "Theme color")["explicit_text_styles"][0]
        base, *modifiers = style["color"].split(";")
        self.assertEqual(base, "theme:accent1")
        self.assertEqual(set(modifiers), {"lumMod=75000", "alpha=50000"})

    def test_attribute_based_colors_are_preserved_without_merging(self):
        cases = (
            ("hslClr", ({"hue": "0", "sat": "50000", "lum": "50000"},
                        {"hue": "60000", "sat": "50000", "lum": "50000"})),
            ("scrgbClr", ({"r": "100000", "g": "0", "b": "0"},
                          {"r": "0", "g": "100000", "b": "0"})),
        )
        for kind, colors in cases:
            shape = add_text(self.slide.shapes, kind, "First")
            shape.text_frame.paragraphs[0].add_run().text = "Second"
            for run, attributes in zip(shape.text_frame.paragraphs[0].runs, colors):
                fill = OxmlElement("a:solidFill")
                color = OxmlElement("a:" + kind)
                for key, value in attributes.items():
                    color.set(key, value)
                fill.append(color)
                run._r.get_or_add_rPr().append(fill)
        report = self.inspect([1])
        for kind, expected_colors in cases:
            with self.subTest(kind=kind):
                styles = styles_for(report, kind)["explicit_text_styles"]
                self.assertEqual(len(styles), 2)
                actual = []
                for style in styles:
                    actual_kind, raw_attributes = style["color"].split(":", 1)
                    self.assertEqual(actual_kind, kind)
                    actual.append(dict(item.split("=", 1) for item in raw_attributes.split(",")))
                self.assertCountEqual(actual, expected_colors)

    def test_style_helpers_do_not_mutate_in_memory_xml(self):
        shape = add_text(self.slide.shapes, "Pure read", "No explicit format")
        before = shape._element.xml
        INSPECTOR.font_style(None)
        INSPECTOR.frame_styles(shape.text_frame)
        self.assertEqual(shape._element.xml, before)
        run = shape.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(20)
        properties = run._r.rPr
        before = shape._element.xml
        style = INSPECTOR.font_style(properties)
        counts = INSPECTOR.frame_styles(shape.text_frame)
        self.assertEqual(style["size"], "20pt")
        self.assertTrue(counts)
        self.assertEqual(shape._element.xml, before)

    def test_paragraph_default_does_not_masquerade_as_run_style(self):
        shape = add_text(self.slide.shapes, "Inherited", "Inherited text")
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.font.name = "Calibri"
        paragraph.font.size = Pt(24)
        paragraph.font.bold = True
        element = styles_for(self.inspect([1]), "Inherited")
        run_style = element["explicit_text_styles"][0]
        self.assertEqual(run_style["font"], "<inherited>")
        self.assertEqual(run_style["size"], "<inherited>")
        self.assertEqual(run_style["weight"], "<inherited>")
        default = element["paragraph_defaults"][0]
        self.assertEqual((default["paragraph"], default["level"]), (1, 0))
        self.assertEqual(default["explicit_text_style"]["size"], "24pt")
        self.assertEqual(default["explicit_text_style"]["font"], "Calibri")

    def test_placeholder_inheritance_stays_unknown(self):
        slide = self.deck.slides.add_slide(self.deck.slide_layouts[1])
        slide.shapes.title.text = "Inherited title"
        report = self.inspect([2])
        title = styles_for(report, slide.shapes.title.name)
        style = title["explicit_text_styles"][0]
        self.assertEqual(style["size"], "<inherited>")
        self.assertEqual(style["font"], "<inherited>")

    def test_group_descendants_are_counted_and_transformed(self):
        group = self.slide.shapes.add_group_shape()
        group.name = "Group"
        child = add_text(group.shapes, "Child", "Grouped text", left=1, top=2, width=2)
        child.text_frame.paragraphs[0].runs[0].font.size = Pt(19)
        group.left, group.top = Inches(4), Inches(3)
        group.width, group.height = Inches(4), Inches(2)
        report = self.inspect([1])
        element = styles_for(report, "Child")
        self.assertEqual(element["coordinate_status"], "slide")
        self.assertEqual(element["box"], {
            "left_pct": 40.0, "top_pct": 40.0,
            "width_pct": 40.0, "height_pct": 26.67,
        })
        self.assertEqual(element["path"], [group.shape_id, child.shape_id])
        self.assertEqual(element["parent_path"], [group.shape_id])
        self.assertEqual(report["slides"][0]["text_boxes"], 1)

    def test_flipped_group_maps_child_to_opposite_side(self):
        group = self.slide.shapes.add_group_shape()
        add_text(group.shapes, "Left child", "Left", left=1, top=2, width=1)
        add_text(group.shapes, "Right child", "Right", left=3, top=2, width=1)
        group.left, group.top = Inches(2), Inches(1)
        group.width, group.height = Inches(6), Inches(2)
        group._element.xfrm.set("flipH", "1")
        report = self.inspect([1])
        left = styles_for(report, "Left child")["box"]
        right = styles_for(report, "Right child")["box"]
        self.assertEqual((left["left_pct"], right["left_pct"]), (60.0, 20.0))
        self.assertEqual(left["width_pct"], 20.0)

    def test_nested_group_paths_are_unique(self):
        outer = self.slide.shapes.add_group_shape()
        inner = outer.shapes.add_group_shape()
        child = add_text(inner.shapes, "Nested", "Nested text")
        element = styles_for(self.inspect([1]), "Nested")
        self.assertEqual(element["path"], [outer.shape_id, inner.shape_id, child.shape_id])
        self.assertEqual(element["parent_path"], [outer.shape_id, inner.shape_id])

    def test_rotated_group_does_not_report_false_precision(self):
        group = self.slide.shapes.add_group_shape()
        add_text(group.shapes, "Rotated", "Rotated group text")
        group.rotation = 30
        element = styles_for(self.inspect([1]), "Rotated")
        self.assertEqual(element["coordinate_status"], "unsupported_group_rotation")
        self.assertIsNone(element["box"])
        self.assertEqual(element["region"], "unresolved")
        self.assertEqual(element["text"], "Rotated group text")

    def test_leaf_rotation_keeps_nominal_box_and_angle(self):
        shape = add_text(self.slide.shapes, "Rotated leaf", "Text")
        shape.rotation = 15
        element = styles_for(self.inspect([1]), "Rotated leaf")
        self.assertEqual(element["coordinate_status"], "slide")
        self.assertEqual(element["rotation_degrees"], 15)
        self.assertEqual(element["box"]["left_pct"], 10)

    def test_degenerate_group_does_not_divide_by_zero(self):
        group = self.slide.shapes.add_group_shape()
        add_text(group.shapes, "Degenerate", "Still readable")
        group._element.xfrm.chExt.cx = 0
        element = styles_for(self.inspect([1]), "Degenerate")
        self.assertEqual(element["coordinate_status"], "unsupported_group_extent")
        self.assertIsNone(element["box"])
        self.assertEqual(element["text"], "Still readable")

    def test_picture_inside_group_is_included_in_image_count(self):
        # One transparent pixel is sufficient to exercise picture traversal.
        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aL1kAAAAASUVORK5CYII="
        )
        group = self.slide.shapes.add_group_shape()
        picture = group.shapes.add_picture(BytesIO(png), Inches(1), Inches(2))
        picture.name = "Grouped picture"
        report = self.inspect([1])
        element = styles_for(report, "Grouped picture")
        self.assertEqual(report["slides"][0]["images"], 1)
        self.assertEqual(element["path"], [group.shape_id, picture.shape_id])

    def test_table_cell_text_and_style_are_exposed(self):
        shape = self.slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(2))
        shape.name = "Evidence table"
        cell = shape.table.cell(1, 0)
        cell.text = "0.25 mm"
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.name = "Arial"
        run.font.size = Pt(16)
        element = styles_for(self.inspect([1]), "Evidence table")
        found = next(cell for cell in element["table_cells"] if cell["text"] == "0.25 mm")
        self.assertEqual((found["row"], found["column"]), (2, 1))
        self.assertEqual(found["explicit_text_styles"][0]["size"], "16pt")

    def test_merged_table_cells_do_not_duplicate_aggregated_text(self):
        shape = self.slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(6), Inches(1))
        shape.name = "Merged table"
        left, right = shape.table.cell(0, 0), shape.table.cell(0, 1)
        left.text, right.text = "Left", "Right"
        left.merge(right)
        # Spanned XML can still contain text even though PowerPoint displays only
        # the merge origin; it must not be counted again as visible evidence.
        right.text = "Hidden spanned cell"
        right.text_frame.paragraphs[0].runs[0].font.name = "HiddenFont"
        report = self.inspect([1])
        counts = [style for styles in report["explicit_text_styles_by_region"].values()
                  for style in styles]
        self.assertEqual(sum(style["characters"] for style in counts), 9)
        self.assertNotIn("HiddenFont", {style["font"] for style in counts})
        cells = styles_for(report, "Merged table")["table_cells"]
        self.assertTrue(cells[0]["is_merge_origin"])
        self.assertTrue(cells[1]["is_spanned"])

    def test_selection_limits_styles_but_preserves_all_summaries(self):
        first = add_text(self.slide.shapes, "First", "First")
        first.text_frame.paragraphs[0].runs[0].font.name = "Arial"
        second_slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])
        second = add_text(second_slide.shapes, "Second", "Second")
        second.text_frame.paragraphs[0].runs[0].font.name = "Calibri"
        report = self.inspect([2])
        self.assertEqual(report["slide_count"], 2)
        self.assertEqual(len(report["slides"]), 2)
        self.assertEqual([detail["slide"] for detail in report["selected_slide_details"]], [2])
        fonts = {style["font"]
                 for styles in report["explicit_text_styles_by_region"].values()
                 for style in styles}
        self.assertIn("Calibri", fonts)
        self.assertNotIn("Arial", fonts)

    def test_detail_order_matches_requested_order(self):
        self.deck.slides.add_slide(self.deck.slide_layouts[6])
        self.assertEqual([entry["slide"] for entry in self.inspect([2, 1])["selected_slide_details"]],
                         [2, 1])

    def test_omitted_selection_has_summary_without_element_details(self):
        add_text(self.slide.shapes, "Summary", "Summary text")
        report = self.inspect()
        self.assertEqual(report["selected_slide_details"], [])
        self.assertEqual(report["slides"][0]["text_boxes"], 1)
        self.assertEqual(report["slide_size_inches"], [10.0, 7.5])

    def test_inspection_never_changes_source(self):
        add_text(self.slide.shapes, "Read only", "Do not modify")
        self.deck.save(self.path)
        original_hash = hashlib.sha256(self.path.read_bytes()).hexdigest()
        INSPECTOR.inspect_presentation(self.path, slides=[1])
        result = self.cli(self.path, "--slides", "1")
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8", errors="replace"))
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), original_hash)

    def test_cli_json_is_utf8_even_with_ascii_environment(self):
        add_text(self.slide.shapes, "中文框", "中文测量结论")
        self.deck.save(self.path)
        result = self.cli(self.path, "--slides", "1", encoding="ascii")
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8", errors="replace"))
        report = json.loads(result.stdout.decode("utf-8"))
        self.assertEqual(styles_for(report, "中文框")["text"], "中文测量结论")

    def test_cli_output_file_is_utf8_json(self):
        add_text(self.slide.shapes, "中文框", "中文测量结论")
        self.deck.save(self.path)
        output = Path(self.temp.name) / "report.json"
        result = self.cli(self.path, "--slides", "1", "--output", output, encoding="ascii")
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8", errors="replace"))
        self.assertEqual(result.stdout, b"")
        report = json.loads(output.read_bytes().decode("utf-8"))
        self.assertEqual(styles_for(report, "中文框")["text"], "中文测量结论")

    def test_cli_output_cannot_overwrite_source(self):
        self.deck.save(self.path)
        original_hash = hashlib.sha256(self.path.read_bytes()).hexdigest()
        result = self.cli(self.path, "--output", self.path)
        self.assertNotEqual(result.returncode, 0)
        self.assertNotIn(b"Traceback", result.stderr)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), original_hash)

    def test_cli_output_cannot_overwrite_hard_linked_source(self):
        self.deck.save(self.path)
        alias = Path(self.temp.name) / "source-alias.pptx"
        try:
            os.link(self.path, alias)
        except (OSError, NotImplementedError) as error:
            self.skipTest(f"Filesystem cannot create a hard link: {error}")
        original_hash = hashlib.sha256(self.path.read_bytes()).hexdigest()
        result = self.cli(self.path, "--output", alias)
        self.assertNotEqual(result.returncode, 0)
        self.assertNotIn(b"Traceback", result.stderr)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(), original_hash)
        self.assertEqual(hashlib.sha256(alias.read_bytes()).hexdigest(), original_hash)

    def test_cli_errors_are_concise_and_do_not_show_tracebacks(self):
        self.deck.save(self.path)
        corrupt = Path(self.temp.name) / "corrupt.pptx"
        corrupt.write_bytes(b"This is not a PowerPoint archive")
        missing = Path(self.temp.name) / "missing.pptx"
        for arguments in ((self.path, "--slides", "bad"),
                          (self.path, "--slides", "2"), (corrupt,), (missing,)):
            with self.subTest(arguments=arguments):
                result = self.cli(*arguments)
                self.assertNotEqual(result.returncode, 0)
                self.assertEqual(result.stdout, b"")
                self.assertTrue(result.stderr)
                self.assertNotIn(b"Traceback", result.stderr)


if __name__ == "__main__":
    unittest.main()
